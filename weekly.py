#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
债市周报 CLI 薄壳
==================
只做三件事：定位输入文件 → 识别 Part/周期 → 调 engine.generate_report()。
核心逻辑在 engine.py（无状态纯任务接口），CLI / 双击脚本 / 未来 Web 都是壳。

用法：
  python3 weekly.py                              # 全自动（债券更新目录取最新文件）
  python3 weekly.py --template X --newbond Y --nafmii Z [--start --end]
  python3 weekly.py --replay-cutoff "2026-08-15 23:59:59"   # 快照重放
  python3 weekly.py --skip-bse                              # 可生成预览（若该 Part 要求 BSE）
"""

import argparse
import datetime
import json
import logging
import os
import re
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.core import engine as core

APP_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(APP_DIR, "output")     # 交付区（复制最近一次结果）
JOBS_DIR = os.path.join(APP_DIR, "jobs")         # 每次任务独立 workspace
SNAPSHOT_DB = os.path.join(APP_DIR, "data", "exchange_snapshots.db")
PARTS_PATH = os.path.join(APP_DIR, "app", "core", "parts.yaml")
DEFAULT_DATADIR = os.path.expanduser("~/Desktop/债券更新")

log = logging.getLogger("weekly")
logging.basicConfig(
    level=logging.INFO,
    format="[%(asctime)s] %(levelname)s - %(message)s",
    handlers=[
        logging.FileHandler(os.path.join(APP_DIR, "运行日志_weekly.log"),
                            mode="a", encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(JOBS_DIR, exist_ok=True)

_DATE_IN_NAME_RE = re.compile(r"(20\d{6})")
_TITLE_RANGE_RE = re.compile(r"[\(（]\s*(\d{1,2})\.(\d{1,2})\s*-\s*(\d{1,2})\.(\d{1,2})\s*[\)）]")


# ---------------------------------------------------------------- 文件定位

def date_from_filename(path):
    m = _DATE_IN_NAME_RE.search(os.path.basename(path))
    if not m:
        return None
    s = m.group(1)
    return datetime.date(int(s[:4]), int(s[4:6]), int(s[6:8]))


def pick_latest(datadir, pattern, keyword=None):
    cands = []
    if not os.path.isdir(datadir):
        return None
    for f in os.listdir(datadir):
        if f.startswith("~$") or not f.endswith(".xlsx"):
            continue
        if not pattern.match(f):
            continue
        if keyword and keyword not in f:
            continue
        cands.append(os.path.join(datadir, f))
    if not cands:
        return None
    cands.sort(key=lambda p: (date_from_filename(p) or datetime.date(1900, 1, 1),
                              os.path.getmtime(p)), reverse=True)
    return cands[0]


# ---------------------------------------------------------------- 周期识别

def parse_prev_week(template_path):
    """从主 sheet A1 的实际 shared-string 引用解析上一期报告区间。"""
    from app.core import template_engine as te
    fname = os.path.basename(template_path)
    m_name = _DATE_IN_NAME_RE.search(fname)
    file_date = m_name.group(1) if m_name else None
    entries = te.read_zip(template_path)
    sheet_file = te.find_main_sheet(entries)[1]
    sheet_xml = entries[sheet_file].decode("utf-8")
    m_a1 = re.search(r'<c r="A1"[^>]*t="s"[^>]*><v>(\d+)</v>', sheet_xml)
    if not m_a1:
        raise SystemExit("无法从主 sheet 定位 A1 的 shared-string 引用")
    si_idx = int(m_a1.group(1))
    ss = entries["xl/sharedStrings.xml"].decode("utf-8")
    sis = re.findall(r"<si>.*?</si>|<si[^>]*/>", ss, re.S)
    if si_idx >= len(sis):
        raise SystemExit(f"A1 引用的 shared-string 索引 {si_idx} 越界（共 {len(sis)}）")
    m_title = _TITLE_RANGE_RE.search(sis[si_idx])
    if not m_title:
        raise SystemExit("无法从 A1 标题解析上一期报告区间（格式应为 (MM.DD-MM.DD)）")
    year = int(file_date[:4]) if file_date else datetime.date.today().year
    prev_start = datetime.date(year, int(m_title.group(1)), int(m_title.group(2)))
    prev_end = datetime.date(year, int(m_title.group(3)), int(m_title.group(4)))
    if prev_end < prev_start:
        prev_start = prev_start.replace(year=year - 1)
    return prev_start, prev_end


def next_week_range(prev_start, prev_end, data_file_dates, week_end=5):
    # 周期 = 上期模板 + week_end + 法定节假日（与网页版同一套逻辑）；
    # 数据文件日期仅在“数据日更早”时收紧（源文件提前提供）
    from app.web.period import next_week_range as _nwr
    start, aligned = _nwr(prev_start, prev_end, week_end=week_end)
    end = aligned
    for d in data_file_dates:
        if d and d < end:
            end = d
    if end <= start:
        end = aligned
    return start, end


# ---------------------------------------------------------------- 回读校验

def find_manifest_for(output_path):
    """在 jobs/*/manifest.json 中按输出文件名找生成存档（取 generated_at 最新者）。"""
    target = os.path.basename(output_path)
    best, best_ts = None, ""
    for d in os.listdir(JOBS_DIR):
        mf = os.path.join(JOBS_DIR, d, "manifest.json")
        if not os.path.isfile(mf):
            continue
        try:
            with open(mf, encoding="utf-8") as f:
                m = json.load(f)
        except (OSError, ValueError):
            continue
        if os.path.basename(m.get("output_file") or "") == target:
            if (m.get("generated_at") or "") >= best_ts and m.get("generated_texts"):
                best, best_ts = mf, m.get("generated_at") or ""
    return best


def _first_diff(expected, actual):
    """返回第一个差异点的位置与上下文（单行摘要）。"""
    exp_lines, act_lines = (expected or "").split("\n"), (actual or "").split("\n")
    for i in range(max(len(exp_lines), len(act_lines))):
        e = exp_lines[i] if i < len(exp_lines) else "<缺失>"
        a = act_lines[i] if i < len(act_lines) else "<缺失>"
        if e != a:
            return f"第{i + 1}行:\n    生成时: {e[:80]}\n    当前值: {a[:80]}"
    return "未知差异"


def verify_file(target_path, manifest_path=None):
    """回读校验：xlsx 的 A1 标题与第 3 行 A-D 文本 vs manifest.generated_texts。"""
    import json as _json
    import openpyxl
    if not os.path.isfile(target_path):
        raise SystemExit(f"文件不存在：{target_path}")
    mf = manifest_path or find_manifest_for(target_path)
    if not mf:
        raise SystemExit(
            f"在 jobs/ 中找不到 {os.path.basename(target_path)} 的生成存档，"
            f"请用 --manifest 显式指定 manifest.json")
    with open(mf, encoding="utf-8") as f:
        m = _json.load(f)
    texts = m.get("generated_texts") or {}
    if not texts:
        raise SystemExit(f"{mf} 中没有 generated_texts 存档（旧版本任务无法回读校验）")

    wb = openpyxl.load_workbook(target_path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    actual = {
        "title": str(ws.cell(row=1, column=1).value or ""),
        "a3": str(ws.cell(row=3, column=1).value or ""),
        "b3": str(ws.cell(row=3, column=2).value or ""),
        "c3": str(ws.cell(row=3, column=3).value or ""),
        "d3": (str(ws.cell(row=3, column=4).value or "")
               if texts.get("d3") is not None else None),
    }
    wb.close()

    print("=" * 70)
    print(f"回读校验: {os.path.basename(target_path)}")
    print(f"生成存档: {mf}（{m.get('generated_at', '')}，{m.get('part_name', '')}）")
    print("-" * 70)
    problems = []
    for key, label in (("title", "A1 标题"), ("a3", "A3 一周回顾"),
                       ("b3", "B3 公司债项目"), ("c3", "C3 交易商协会"),
                       ("d3", "D3 企业债")):
        if actual.get(key) is None:
            continue
        if actual[key] != (texts.get(key) or ""):
            problems.append(label)
            print(f"[差异] {label}")
            print("  " + _first_diff(texts.get(key) or "", actual[key]))
    if problems:
        print("-" * 70)
        print(f"结论: 生成后内容被改动 → {', '.join(problems)}（请人工确认是否为有意修改）")
        raise SystemExit(2)
    print("结论: 与生成时存档逐字一致（生成后无改动或改动已还原）")
    print("=" * 70)


# ---------------------------------------------------------------- 主流程

def main():
    parser = argparse.ArgumentParser(description="债市周报（Mac 版，核心引擎薄壳）")
    parser.add_argument("--template", default=None, help="上周周报 xlsx（缺省自动找最新）")
    parser.add_argument("--newbond", default=None, help="本周新发行债券 xlsx（缺省自动找最新）")
    parser.add_argument("--nafmii", default=None, help="本周 NAFMII 进度 xlsx（缺省自动找最新）")
    parser.add_argument("--datadir", default=DEFAULT_DATADIR, help="数据目录（自动定位时用）")
    parser.add_argument("--map", default=os.path.join(APP_DIR, "input", "简称.xlsx"))
    parser.add_argument("--start", default=None, help="覆盖本期开始日期 yyyy-mm-dd")
    parser.add_argument("--end", default=None, help="覆盖本期结束日期 yyyy-mm-dd")
    parser.add_argument("--skip-bse", action="store_true", help="跳过北交所（requires_bse 时产出预览版）")
    parser.add_argument("--bse-retries", type=int, default=2,
                        help="北交所抓取失败后的自动重试次数（默认 2，即共 3 次尝试；0 = 失败即停）")
    parser.add_argument("--bse-retry-wait", type=int, default=600,
                        help="北交所重试间隔秒数（默认 600；WAF 临时拦截通常几分钟内自行解除）")
    parser.add_argument("--nafmii-resolve", action="append", default=[],
                        help="兼容旧版参数；当前同日多状态按审核流程取更靠后的状态")
    parser.add_argument("--nafmii-force-first", action="store_true",
                        help="兼容旧版参数；当前同日多状态按审核流程取更靠后的状态")
    parser.add_argument("--replay-cutoff", default=None,
                        help="从快照库重放（不抓官网），截止时点 yyyy-mm-dd[ hh:mm:ss]")
    parser.add_argument("--out", default=None, help="额外复制结果到指定路径")
    parser.add_argument("--verify", default=None, metavar="XLSX",
                        help="回读校验：把最终 xlsx 的标题与 A3-D3 文本和生成时存档比对，"
                             "列出生成后被人工改动的内容（防手误）")
    parser.add_argument("--manifest", default=None, metavar="JSON",
                        help="配合 --verify：显式指定 manifest.json（缺省在 jobs/ 里按输出文件名自动查找）")
    args = parser.parse_args()

    if args.verify:
        verify_file(os.path.abspath(args.verify),
                    os.path.abspath(args.manifest) if args.manifest else None)
        return

    datadir = os.path.expanduser(args.datadir)
    template_path = args.template or pick_latest(
        datadir, re.compile(r"^.*(债市周报|债券周报).*\.xlsx$"))
    newbond_path = args.newbond or pick_latest(datadir, re.compile(r"^.*新发行债券.*\.xlsx$"))
    nafmii_path = args.nafmii or pick_latest(datadir, re.compile(r"^.*NAFMII.*\.xlsx$"))
    if not template_path:
        raise SystemExit(f"未指定 --template，且在 {datadir} 找不到周报模板文件")
    for p, tag in ((newbond_path, "新发行债券"), (nafmii_path, "NAFMII")):
        if not p:
            raise SystemExit(f"未指定 --{tag}，且在 {datadir} 找不到对应输入文件")
    for p in (template_path, newbond_path, nafmii_path):
        if not os.path.exists(p):
            raise SystemExit(f"文件不存在：{p}")
    log.info("输入文件：模板=%s 新发行=%s NAFMII=%s", template_path, newbond_path, nafmii_path)

    # ---- 识别 Part 与周期 ----
    parts, _ = core.load_parts(PARTS_PATH)
    part = core.identify_part(os.path.basename(template_path), parts)
    if part is None:
        raise SystemExit(
            f"无法从文件名「{os.path.basename(template_path)}」识别地区 Part，"
            f"可用关键词：{[kw for p in parts for kw in p['keywords']]}")
    rules = core.part_rules(part, core.load_parts(PARTS_PATH)[1])
    prev_start, prev_end = parse_prev_week(template_path)
    from app.web import period as web_period
    week_end = web_period.resolve_week_end(rules, prev_end)
    start, end = next_week_range(prev_start, prev_end,
                                 [date_from_filename(newbond_path),
                                  date_from_filename(nafmii_path)],
                                 week_end=week_end)
    if args.start:
        start = datetime.datetime.strptime(args.start, "%Y-%m-%d").date()
    if args.end:
        end = datetime.datetime.strptime(args.end, "%Y-%m-%d").date()
    if end <= start:
        raise SystemExit(f"日期区间非法：{start} ~ {end}")
    print(f"已识别：{part['name']}")
    print(f"包含地区：{'、'.join(part['provinces'])}")
    print(f"上期报告：{prev_start:%m.%d}–{prev_end:%m.%d}")
    print(f"本期报告：{start:%m.%d}–{end:%m.%d}")

    # ---- 构造任务并执行 ----
    job_id = core.new_job_id()
    ctx = core.RunContext(
        job_id=job_id,
        workspace=os.path.join(JOBS_DIR, job_id),
        snapshot_db=SNAPSHOT_DB,
        parts_path=PARTS_PATH,
        logger=log,
    )
    # 旧版人工裁决参数解析（仅为兼容旧调用；Core 当前固定按审核流程取更靠后的状态）
    manual_resolutions = []
    for item in args.nafmii_resolve:
        if "=" not in item:
            raise SystemExit(f"--nafmii-resolve 格式应为 项目名=状态（收到：{item}）")
        proj_part, chosen = item.split("=", 1)
        proj, date = proj_part, None
        if "|" in proj_part:
            proj, date = proj_part.split("|", 1)
        manual_resolutions.append({"project": proj.strip(), "date": date, "chosen": chosen.strip()})

    req = core.ReportRequest(
        part_id=part["id"],
        start=start.strftime("%Y-%m-%d"),
        end=end.strftime("%Y-%m-%d"),
        template_path=os.path.abspath(template_path),
        newbond_path=os.path.abspath(newbond_path),
        nafmii_path=os.path.abspath(nafmii_path),
        map_path=os.path.abspath(args.map),
        exchange_mode="replay" if args.replay_cutoff else "live",
        replay_cutoff=args.replay_cutoff,
        skip_bse=args.skip_bse,
        bse_retries=max(0, args.bse_retries),
        bse_retry_wait=max(0, args.bse_retry_wait),
        manual_resolutions=manual_resolutions,
        force_first_all=args.nafmii_force_first,
    )
    result = core.generate_report(req, ctx)

    # ---- 复制到交付区并打印 QA ----
    final_name = os.path.basename(result.output_file)
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    shutil.copy2(result.output_file, os.path.join(OUTPUT_DIR, final_name))
    if args.out:
        # 交付纪律：只有 PASS 才允许复制为调用者指定的“最终版”文件名；
        # WARNING/BLOCKED 一律拒绝，防止预览版被改名成误导性的最终名。
        if result.status == "PASS":
            shutil.copy2(result.output_file, args.out)
        else:
            log.warning("状态 %s：拒绝把预览版复制为指定的最终文件名 %s",
                        result.status, args.out)

    print()
    print("=" * 70)
    print(f"QA 状态: {result.status}")
    print(f"  SSE: {result.qa['sse']}   SZSE: {result.qa['szse']}   "
          f"BSE: {result.qa['bse']}   NAFMII: {result.qa['nafmii']}   "
          f"OOXML: {result.qa['ooxml']}")
    for k, v in result.metrics.items():
        print(f"  {k}: {v}")
    for w in result.warnings:
        print(f"  [WARNING] {w}")
    for b in result.blockers:
        label = b.get('project') or '/'.join(str(b.get(k, '')) for k in ('source', 'province', 'project_id'))
        detail = ' / '.join(b.get('candidates') or b.get('dates') or [])
        print(f"  [BLOCKED] {label[:80]}（{detail}）")
    for r in result.manual_resolutions:
        print(f"  [自动选取] {r['project'][:40]}（{r['date']}: "
              f"{' → '.join(r['candidates'])} → 采用流程后位状态 {r['chosen']}，方式 {r['method']}）")
    print("-" * 70)
    print(f"输出: {result.output_file}")
    print(f"交付区副本: {os.path.join(OUTPUT_DIR, final_name)}")
    if result.status != "PASS":
        print(f"注意: 状态为 {result.status}，本次产物为预览版，不可直接交付")
    print("=" * 70)


if __name__ == "__main__":
    main()
